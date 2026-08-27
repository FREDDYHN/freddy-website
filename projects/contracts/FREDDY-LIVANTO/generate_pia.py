# -*- coding: utf-8 -*-
"""生成《个人信息保护影响评估报告》(PIA) —— FREDDY → LIVANTO，供 PIPL 标准合同备案。"""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(OUT_DIR, '个人信息保护影响评估报告.docx')

FONT = '宋体'
FONT_EN = 'Times New Roman'

# (标题, [段落...])；标题空字符串表示不单独加标题
SECTIONS = [
    ('个人信息保护影响评估报告', []),
    ('', ['评估主体：福瑞笛（上海）信息咨询有限公司淮南分公司（FREDDY (SHANGHAI) INFORMATION CONSULTING LTD. HN）',
          '评估日期：____年____月____日']),
    ('一、评估背景与依据', [
        '依据《中华人民共和国个人信息保护法》第三十八条、第五十五条以及《个人信息出境标准合同办法》第五条，'
        '福瑞笛（上海）信息咨询有限公司淮南分公司（以下简称"福瑞笛"）就其向境外接收方 LIVANTO GmbH（以下简称"LIVANTO"）'
        '提供个人信息开展个人信息保护影响评估。',
    ]),
    ('二、个人信息出境基本情况', [
        '（一）个人信息处理者（出境方）',
        '名称：福瑞笛（上海）信息咨询有限公司淮南分公司（FREDDY (SHANGHAI) INFORMATION CONSULTING LTD. HN）',
        '地址：安徽省淮南市中环国际广场158金融中心四层418室',
        '统一社会信用代码：91340400MADDK97K4X',
        '负责人：桂程程',
        '',
        '（二）境外接收方',
        '名称：LIVANTO GmbH',
        '地址：德国多特蒙德 Luisenhoffnung 3C（邮编 44225）',
        '商业登记号：HRB 38628',
        '增值税识别号：DE 464031041',
        '',
        '（三）出境个人信息的情况',
        '个人信息主体类别：客户（中国制造商）的联系人及法定代表人',
        '个人信息种类：姓名、公司名称、地址、电子邮箱、电话号码、微信号、税号、包装数据',
        '敏感个人信息：无',
        '处理目的：履行德国《包装法》授权代表服务（LUCID 注册、双元系统参与、数据申报、完整性声明）',
        '处理方式：个人信息的电子化记录、存储与传输',
        '传输频率：持续、按需',
        '保存期限：客户授权代表合同存续期间及法定保存义务期间',
        '保存地点：德国',
        '出境规模：自上年1月1日起累计向境外提供个人信息不满 10 万人',
    ]),
    ('三、合法性、正当性、必要性评估', [
        '（一）合法性：福瑞笛基于与客户签订的授权代表合同收集个人信息，并依据《个人信息保护法》第三十九条取得个人信息主体的单独同意；'
        '向境外提供个人信息依据《个人信息保护法》第三十八条，通过与境外接收方订立《个人信息出境标准合同》的方式进行。',
        '（二）正当性：向境外提供个人信息是为客户履行德国包装法合规义务所必需，处理目的正当、明确。',
        '（三）必要性：出境个人信息限于实现授权代表服务目的所需的最小范围，未超出必要限度。',
    ]),
    ('四、出境风险及对个人权益的影响评估', [
        '（一）出境个人信息均为一般个人信息（不含敏感个人信息），规模小、敏感程度低。',
        '（二）可能存在的风险：跨境传输过程中数据被篡改、破坏、泄露、丢失、非法利用的风险；个人信息主体维权不便的风险。',
        '（三）综合风险等级：低。',
    ]),
    ('五、境外接收方保障能力评估', [
        '（一）LIVANTO 为德国注册的有限责任公司，受欧盟《通用数据保护条例》（GDPR）约束，所在地区个人信息保护水平较高。',
        '（二）LIVANTO 承诺依据标准合同条款处理个人信息，采取数据传输加密（TLS）、访问与权限控制、最小授权、删除方案、防止未经授权访问等技术与组织措施。',
        '（三）双方已订立《个人信息出境标准合同》，明确境外接收方的义务及个人信息主体的第三方受益权。',
    ]),
    ('六、出境后安全风险及维权渠道评估', [
        '（一）LIVANTO 采取加密传输、访问控制等措施，有效降低出境个人信息被篡改、破坏、泄露、丢失、非法利用的风险。',
        '（二）个人信息主体可依据标准合同约定的第三方受益权，向福瑞笛或 LIVANTO 主张权利，亦可向监管机构投诉或向有管辖权的法院提起诉讼，维权渠道畅通。',
    ]),
    ('七、境外接收方所在国法律政策影响评估', [
        '（一）德国作为欧盟成员国，适用 GDPR，赋予个人信息主体充分的权利，个人信息保护水平高。',
        '（二）德国现行法律政策不影响 LIVANTO 履行标准合同义务；如相关法律政策发生变化并可能影响履约，LIVANTO 应立即通知福瑞笛。',
    ]),
    ('八、其他事项', ['无其他可能影响个人信息出境安全的事项。']),
    ('九、评估结论', [
        '经评估，福瑞笛向 LIVANTO 提供个人信息的目的、范围、方式合法、正当、必要；出境个人信息规模小、敏感程度低；'
        '境外接收方具备充分的保障能力；出境风险可控。福瑞笛可以依据《个人信息保护法》第三十八条，'
        '通过与 LIVANTO 订立《个人信息出境标准合同》的方式向境外提供个人信息，并依规定向所在地省级网信部门备案。',
    ]),
    ('', ['评估人（签字）：____________________', '日期：____年____月____日']),
]


def _set_run(run, size=Pt(12), bold=False):
    run.font.name = FONT_EN
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:ascii'), FONT_EN)
    rfonts.set(qn('w:hAnsi'), FONT_EN)
    rfonts.set(qn('w:eastAsia'), FONT)


def _add_page_number(doc):
    """页脚中央添加页码 (PAGE field)。"""
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run1 = p.add_run()
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    run1._element.append(fld1)
    run2 = p.add_run()
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    run2._element.append(instr)
    run3 = p.add_run()
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run3._element.append(fld2)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Mm(25)
    sec.right_margin = Mm(25)
    sec.top_margin = Mm(25)
    sec.bottom_margin = Mm(25)

    style = doc.styles['Normal']
    style.font.name = FONT_EN
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    first_title = True
    for title, paras in SECTIONS:
        if title:
            p = doc.add_paragraph()
            if first_title:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_run(p.add_run(title), size=Pt(16), bold=True)
                first_title = False
            else:
                _set_run(p.add_run(title), size=Pt(13), bold=True)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        for text in paras:
            p = doc.add_paragraph()
            if text == '':
                p.paragraph_format.space_after = Pt(4)
                continue
            _set_run(p.add_run(text), size=Pt(12))
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(2)

    _add_page_number(doc)
    doc.save(OUT_PATH)
    print('Generated:', OUT_PATH)


if __name__ == '__main__':
    build()
