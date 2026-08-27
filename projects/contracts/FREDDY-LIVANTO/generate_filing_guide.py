# -*- coding: utf-8 -*-
"""生成《个人信息出境标准合同备案操作流程》文档（含页码）。"""
import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(OUT_DIR, '个人信息出境标准合同备案操作流程.docx')

FONT = '宋体'
FONT_EN = 'Times New Roman'

SECTIONS = [
    ('个人信息出境标准合同备案操作流程', []),
    ('', ['适用主体：福瑞笛（上海）信息咨询有限公司淮南分公司（FREDDY (SHANGHAI) INFORMATION CONSULTING LTD. HN）']),
    ('一、备案依据', [
        '1.《中华人民共和国个人信息保护法》第三十八条、第三十九条、第五十五条；',
        '2.《个人信息出境标准合同办法》（国家互联网信息办公室令第13号，2023年6月1日起施行）。',
    ]),
    ('二、确认适用「标准合同备案」路径（而非安全评估）', [
        '依据《办法》第四条，福瑞笛同时满足下列情形，适用最轻的「标准合同备案」路径，无需安全评估：',
        '1. 非关键信息基础设施运营者；',
        '2. 处理个人信息不满 100 万人；',
        '3. 自上年1月1日起累计向境外提供个人信息不满 10 万人；',
        '4. 自上年1月1日起累计向境外提供敏感个人信息不满 1 万人。',
    ]),
    ('三、备案前准备（三步）', [
        '1. 开展个人信息保护影响评估（PIA）—— 已备《个人信息保护影响评估报告》；',
        '2. 签署《个人信息出境标准合同》—— 已备（FREDDY 为出境方、LIVANTO GmbH 为境外接收方）；',
        '3. 取得个人信息主体单独同意 —— 已写入客户授权代表合同 §10(3)（客户签约时同意将信息提供给 LIVANTO）。',
    ]),
    ('四、备案提交（《办法》第七条）', [
        '1. 提交时间：标准合同生效之日起 10 个工作日内；',
        '2. 提交对象：福瑞笛所在地省级网信部门 = 安徽省委网信办；',
        '3. 提交材料：',
        '   ①《个人信息出境标准合同》；',
        '   ②《个人信息保护影响评估报告》。',
    ]),
    ('五、备案材料清单（对应「备案材料」文件夹）', [
        '核心材料（必交）：',
        '1. 个人信息出境标准合同_PIPL.pdf',
        '2. 个人信息保护影响评估报告.pdf',
        '备查材料（建议携带）：',
        '3. Kooperationsvereinbarung_FREDDY-LIVANTO.pdf（双方合作协议 + SCC 附件）',
        '4. 包装法_授权代表合同_空白模板.pdf（客户单独同意条款）',
    ]),
    ('六、备案后事项（《办法》第八条）', [
        '出现下列情形之一的，应当重新开展影响评估、补充或重新订立标准合同，并重新履行备案手续：',
        '1. 出境目的、范围、种类、敏感程度、方式、保存地点或保存期限发生变化；',
        '2. 境外接收方所在国家或地区的个人信息保护政策和法规发生变化等可能影响个人信息权益的情形；',
        '3. 其他可能影响个人信息权益的情形。',
    ]),
    ('七、注意事项', [
        '1. 备案为形式审查，材料不齐全或不符合要求将被退回补正；',
        '2. 具体提交渠道（线上「数据出境标准合同备案」系统或线下面交）以安徽省委网信办最新指南为准，建议提前电话或官网确认；',
        '3. 个人信息主体的投诉、举报渠道已在《个人信息出境标准合同》第六条约定；',
        '4. 评估报告中的评估日期、评估人签字需在提交前填写。',
    ]),
]


def _set_run(run, size=Pt(12), bold=False):
    run.font.name = FONT_EN
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:ascii'), FONT_EN)
    rfonts.set(qn('w:hAnsi'), FONT_EN)
    rfonts.set(qn('w:eastAsia'), FONT)


def _add_page_number(doc):
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run1 = p.add_run()
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    run1._element.append(fld1)
    run2 = p.add_run()
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    run2._element.append(instr)
    run3 = p.add_run()
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run3._element.append(fld2)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Mm(25)
    sec.right_margin = Mm(25)
    sec.top_margin = Mm(25)
    sec.bottom_margin = Mm(25)
    style = doc.styles['Normal']
    style.font.name = FONT_EN
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    first = True
    for title, paras in SECTIONS:
        if title:
            p = doc.add_paragraph()
            if first:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_run(p.add_run(title), size=Pt(16), bold=True)
                first = False
            else:
                _set_run(p.add_run(title), size=Pt(13), bold=True)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        for text in paras:
            if text == '':
                continue
            p = doc.add_paragraph()
            _set_run(p.add_run(text), size=Pt(12))
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(2)

    _add_page_number(doc)
    doc.save(OUT_PATH)
    print('Generated:', OUT_PATH)


if __name__ == '__main__':
    build()
